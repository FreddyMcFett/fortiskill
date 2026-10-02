#!/usr/bin/env python3
"""Look up Fortinet SKUs and list prices in a Fortinet price list workbook.

Searches the `DataSet` sheet by its Item, SKU, Description #1 and Price columns.

Usage:
  # free-text search over SKU + Item + Description (case-insensitive, all terms must match)
  python pricelist_lookup.py <pricelist.xlsx> "FortiGate-50G ATP 5 Year"

  # exact/prefix SKU lookup
  python pricelist_lookup.py <pricelist.xlsx> --sku FC-10-GT50G-928-02-60
  python pricelist_lookup.py <pricelist.xlsx> --sku FC-10-GT50G-

  # batch-verify a final BOM: one SKU per line on stdin, reports any that are missing
  cat skus.txt | python pricelist_lookup.py <pricelist.xlsx> --verify

Output: TSV  SKU <tab> Price <tab> Item <tab> Description
Exit code 1 if nothing found / any verify SKU missing.
"""
import sys
import argparse
import warnings


def load_rows(path):
    warnings.filterwarnings("ignore")
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() == "dataset":
            sheet = wb[name]
            break
    if sheet is None:
        sys.exit(f"error: no 'DataSet' sheet in {path} (sheets: {wb.sheetnames[:8]}...)")
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h else "" for h in next(rows)]

    def col(*names):
        for n in names:
            for i, h in enumerate(header):
                if h == n:
                    return i
        return None

    i_sku = col("sku")
    i_price = col("price")
    i_item = col("item")
    i_desc = col("description #1", "description")
    if i_sku is None or i_price is None:
        sys.exit(f"error: DataSet header not recognized: {header}")
    out = []
    for r in rows:
        sku = r[i_sku]
        if not sku:
            continue
        out.append((
            str(sku).strip(),
            r[i_price],
            str(r[i_item]).strip() if i_item is not None and r[i_item] else "",
            str(r[i_desc]).strip() if i_desc is not None and r[i_desc] else "",
        ))
    return out


def emit(matches, limit):
    for sku, price, item, desc in matches[:limit]:
        print(f"{sku}\t{price}\t{item}\t{desc}")
    if len(matches) > limit:
        print(f"... {len(matches) - limit} more matches (raise --limit or refine query)", file=sys.stderr)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pricelist")
    ap.add_argument("query", nargs="?", default=None)
    ap.add_argument("--sku", help="exact SKU, or prefix if it ends with '-'")
    ap.add_argument("--verify", action="store_true", help="read SKUs from stdin, report missing ones")
    ap.add_argument("--limit", type=int, default=25)
    args = ap.parse_args()

    rows = load_rows(args.pricelist)

    if args.verify:
        wanted = [l.strip() for l in sys.stdin if l.strip()]
        index = {sku: (price, desc) for sku, price, _, desc in rows}
        missing = []
        for s in wanted:
            if s in index:
                print(f"OK\t{s}\t{index[s][0]}\t{index[s][1]}")
            else:
                missing.append(s)
                print(f"MISSING\t{s}")
        sys.exit(1 if missing else 0)

    if args.sku:
        q = args.sku.strip()
        if q.endswith("-"):
            matches = [r for r in rows if r[0].upper().startswith(q.upper())]
        else:
            matches = [r for r in rows if r[0].upper() == q.upper()]
    elif args.query:
        terms = args.query.lower().split()
        matches = [r for r in rows if all(t in f"{r[0]} {r[2]} {r[3]}".lower() for t in terms)]
    else:
        ap.error("provide a query, --sku, or --verify")

    if not matches:
        print("no matches", file=sys.stderr)
        sys.exit(1)
    emit(matches, args.limit)


if __name__ == "__main__":
    main()
